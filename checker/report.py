"""Read a flat invoice export CSV, run the readiness rules, write an Excel report."""
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from rules import RULE_TEXT, check, summarise  # noqa: E402

INK, IVORY, BRASS, MUTED, LINE = "0B1220", "F6F1E7", "C9A24C", "6B7280", "E4DCC9"
HEAD_FILL = PatternFill("solid", fgColor=INK)
BAND = PatternFill("solid", fgColor="FBF8F1")
HIGH = PatternFill("solid", fgColor="FDE8E8")
MED = PatternFill("solid", fgColor="FFF4D6")
thin = Side(style="thin", color=LINE)
BORDER = Border(bottom=thin)

SEVERITY = {"IBR-128-AE": "HIGH", "IBR-104-AE": "HIGH", "IBR-002-AE": "HIGH", "IBR-154-AE": "HIGH", "R-04": "HIGH", "R-02": "HIGH",
            "R-11": "HIGH", "R-05": "MEDIUM", "R-10": "MEDIUM", "R-09": "MEDIUM", "R-08": "MEDIUM", "R-03": "MEDIUM", "R-01": "MEDIUM",
            "R-07": "MEDIUM", "R-06": "HIGH"}


def load(path: Path) -> list[dict]:
    with open(path, newline="") as f:
        rows = list(csv.DictReader(f))
    inv = {}
    order = []
    for r in rows:
        key = (r["invoice_no"], r["issue_date"], r["buyer_name"])
        if key not in inv:
            inv[key] = {k: r[k] for k in r if k not in ("line_no", "description", "qty", "unit_price", "line_amount", "vat_category", "vat_rate", "vat_amount", "line_amount_aed", "vat_amount_aed")}
            inv[key]["lines"] = []
            order.append(key)
        inv[key]["lines"].append({k: r[k] for k in ("line_no", "description", "qty", "unit_price", "line_amount", "vat_category", "vat_rate", "vat_amount", "line_amount_aed", "vat_amount_aed")})
    return [inv[k] for k in order]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=IVORY, name="Calibri")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write(invoices, findings, out: Path, client="Sample client"):
    s = summarise(invoices, findings)
    wb = Workbook()
    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"E-invoicing readiness: {client}"
    ws["A1"].font = Font(size=18, bold=True, color=INK, name="Georgia")
    ws["A2"] = f"{s['invoices_clean']} of {s['invoices']} invoices would pass PINT AE validation as exported. {s['invoices_with_findings']} need a fix before an accredited provider can send them."
    ws["A2"].font = Font(size=11, italic=True, color=BRASS)
    ws["A3"] = "Read only: no source file was modified. Rules from PINT AE v1.0.3 and Finamatik readiness checks."
    ws["A3"].font = Font(size=9, italic=True, color=MUTED)
    rows = [("Invoices checked", s["invoices"]), ("Invoice lines checked", s["lines"]), ("Invoices clean as exported", s["invoices_clean"]),
            ("Invoices needing a fix", s["invoices_with_findings"]), ("Readiness", f"{s['readiness_pct']}%"), ("Findings", s["findings"])]
    r = 5
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, color=INK)
        ws.cell(row=r, column=1).fill = BAND
        ws.cell(row=r, column=2, value=v).alignment = Alignment(horizontal="right")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Findings by rule").font = Font(bold=True, size=12, color=INK)
    r += 1
    for c, h in enumerate(["Rule", "What it checks", "Findings", "Severity", "Source"], 1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, r, 5)
    r += 1
    for rule in sorted(s["by_rule"], key=lambda x: -s["by_rule"][x]):
        ws.cell(row=r, column=1, value=rule)
        ws.cell(row=r, column=2, value=RULE_TEXT[rule][0])
        ws.cell(row=r, column=3, value=s["by_rule"][rule]).alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=4, value=SEVERITY.get(rule, "MEDIUM"))
        ws.cell(row=r, column=5, value=RULE_TEXT[rule][1]).font = Font(color=MUTED, size=9)
        if SEVERITY.get(rule) == "HIGH":
            ws.cell(row=r, column=4).fill = HIGH
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="How to read this workbook").font = Font(bold=True, size=12, color=INK)
    for line in ["Exceptions: every finding with the invoice, line, field, value and the fix, ranked HIGH first.",
                 "Fix list: the same findings grouped by the master data field that causes them, which is the actual work.",
                 "Rules: what each check means and where it comes from.",
                 "Clean invoices: the ones that would pass as exported."]:
        r += 1
        ws.cell(row=r, column=1, value=line).font = Font(color=INK)
    autosize(ws, [34, 62, 12, 12, 40])
    # ---- Exceptions
    we = wb.create_sheet("Exceptions")
    we.sheet_view.showGridLines = False
    heads = ["Severity", "Invoice", "Line", "Rule", "Field", "Value as exported", "Finding", "Fix"]
    for c, h in enumerate(heads, 1):
        we.cell(row=1, column=c, value=h)
    style_header(we, 1, len(heads))
    order = {"HIGH": 0, "MEDIUM": 1}
    for i, x in enumerate(sorted(findings, key=lambda x: (order.get(SEVERITY.get(x["rule"], "MEDIUM"), 1), x["invoice_no"], x["line_no"])), 2):
        sev = SEVERITY.get(x["rule"], "MEDIUM")
        vals = [sev, x["invoice_no"], x["line_no"], x["rule"], x["field"], x["value"], x["message"], x["fix"]]
        for c, v in enumerate(vals, 1):
            cell = we.cell(row=i, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (7, 8)))
        we.cell(row=i, column=1).fill = HIGH if sev == "HIGH" else MED
    we.freeze_panes = "A2"
    autosize(we, [10, 16, 6, 12, 26, 20, 52, 56])
    # ---- Fix list grouped by field
    wf = wb.create_sheet("Fix list")
    wf.sheet_view.showGridLines = False
    by_field = defaultdict(list)
    for x in findings:
        by_field[(x["field"], x["fix"])].append(x)
    heads = ["Field to fix", "Where it lives", "Invoices affected", "Findings", "What to do"]
    for c, h in enumerate(heads, 1):
        wf.cell(row=1, column=c, value=h)
    style_header(wf, 1, len(heads))
    WHERE = {"buyer_emirate": "Customer master", "buyer_trn": "Customer master", "seller_trn": "Company settings", "fx_rate": "Invoice / currency table",
             "line_amount_aed / vat_amount_aed": "Invoice lines (export mapping)", "txn_type_code": "Invoice type mapping", "issue_date": "Export format",
             "invoice_no": "Numbering sequence", "vat_category": "Item tax codes", "vat_rate": "Item tax codes", "vat_amount": "Tax calculation setting",
             "line_amount": "Discount handling", "total_excl_vat": "Totals mapping", "total_vat": "Totals mapping", "total_incl_vat": "Totals mapping"}
    r = 2
    for (field, fix), xs in sorted(by_field.items(), key=lambda kv: -len(kv[1])):
        vals = [field, WHERE.get(field, "Invoice"), len({x["invoice_no"] for x in xs}), len(xs), fix]
        for c, v in enumerate(vals, 1):
            cell = wf.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 5))
        r += 1
    autosize(wf, [30, 28, 16, 10, 70])
    # ---- Rules
    wr = wb.create_sheet("Rules")
    wr.sheet_view.showGridLines = False
    for c, h in enumerate(["Rule", "What it checks", "Severity", "Source"], 1):
        wr.cell(row=1, column=c, value=h)
    style_header(wr, 1, 4)
    for i, (rule, (txt, src)) in enumerate(RULE_TEXT.items(), 2):
        for c, v in enumerate([rule, txt, SEVERITY.get(rule, "MEDIUM"), src], 1):
            wr.cell(row=i, column=c, value=v).border = BORDER
    autosize(wr, [14, 70, 12, 40])
    # ---- Clean invoices
    wc = wb.create_sheet("Clean invoices")
    wc.sheet_view.showGridLines = False
    bad = {x["invoice_no"] for x in findings}
    for c, h in enumerate(["Invoice", "Date", "Buyer", "Currency", "Total incl. VAT"], 1):
        wc.cell(row=1, column=c, value=h)
    style_header(wc, 1, 5)
    r = 2
    for inv in invoices:
        if inv["invoice_no"] in bad:
            continue
        for c, v in enumerate([inv["invoice_no"], inv["issue_date"], inv["buyer_name"], inv["currency"], inv["total_incl_vat"]], 1):
            wc.cell(row=r, column=c, value=v).border = BORDER
        r += 1
    autosize(wc, [16, 12, 30, 10, 16])
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)
    return s


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "sample/invoice_export.csv"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "out/einvoicing-readiness.xlsx"
    out.parent.mkdir(exist_ok=True, parents=True)
    invoices = load(src)
    findings = check(invoices)
    s = write(invoices, findings, out)
    print(f"{s['invoices']} invoices, {s['lines']} lines, {s['findings']} findings, readiness {s['readiness_pct']}% -> {out}")
    for k, v in s["by_rule"].most_common():
        print(f"  {k:12} {v}")
